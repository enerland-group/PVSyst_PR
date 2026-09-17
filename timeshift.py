r"""
Desfase horario de PVsyst: T08_time_shifts -> UsesDST de la MEF (o CSV raw
corregido si no es de fiar) + -imt del CLI.

El CLI solo admite `-imt` entre -30 y +30 minutos, asi que el desfase total de la
planta se atiende en dos sitios distintos:

    total (min)  =  horas enteras   ->  cambio de hora (DST)
                 +  resto [-30;+30] ->  -imt del convert-meteo

Las horas enteras son cambios de hora (DST), y la tabla `T08_time_shifts` del
Excel de configuracion dice, por planta, el desfase total en minutos y las FECHAS
de esos cambios (`start_date` y `end_date`).

IMPORTANTE (confirmado con pruebas reales sobre FRA, PL1 y PL2 el 31/07/2026):
cuando la MEF declara `UsesDST=True` + `RefTimeS=LegalTime` + `DateHEte`/
`DateHHiver` coherentes con la regla de la UE (ultimo domingo de marzo/octubre),
PVsyst YA RESUELVE ese cambio de hora el solo dentro de `run-simulation`, sin que
haga falta tocar el CSV. Reescribir las fechas del CSV EN ESE CASO duplica la
correccion (se aplica una vez a mano y otra vez la aplica PVsyst) y descentra el
resultado justo el doble de lo que hacia falta -- o hace que el CLI rechace el
fichero directamente. Por eso, cuando `es_dst_fiable()` dice que la MEF es de
fiar para una fecha de T08, este script NO TOCA el CSV en esa fecha: deja que
PVsyst la resuelva y solo aporta el resto de minutos sueltos por `-imt`.

Cuando la MEF NO es de fiar (UsesDST=False, RefTimeS distinto de LegalTime,
DateHEte/DateHHiver ausentes, del año equivocado o que no cuadran con la regla de
la UE), no hay garantia de que PVsyst vaya a resolver nada, y se cae al mecanismo
antiguo: mirar el CSV raw en la fecha de T08 y corregirlo a mano.

  - si falta una hora  (salto tipo 02:00 -> 03:00, marcha adelante del reloj)
    -> se RESTA una hora a todas las filas siguientes,
  - si hay una hora repetida (el reloj vuelve atras)
    -> se SUMA una hora a todas las filas siguientes,
  - si no hay ni una cosa ni la otra
    -> NO se corrige la hora; solo queda un aviso en el log.

Cuando la fecha del cambio queda FUERA del CSV no hay artefacto que buscar, pero
si hay que actuar (solo en el mecanismo de fallback, sin UsesDST fiable):

  - fecha ANTERIOR al primer dato: el reloj ya habia cambiado, asi que todas las
    filas vienen con el desfase puesto -> se corrige desde la PRIMERA fila, con el
    sentido que da el lado del tramo (inicio = el reloj se adelanto, fin = volvio
    atras). Es el caso de FRA, cuyo CSV empieza el 21/04 con el cambio el 29/03
    (aunque con UsesDST fiable, que es su caso real, esto ya no se aplica).
  - fecha POSTERIOR al ultimo dato: ese cambio no le afecta al fichero.

Con las dos fechas del tramo por detras del CSV las dos correcciones (del
fallback) se suman y se anulan, que es justo lo que toca: el tramo se abrio y se
cerro antes de empezar.

Las correcciones del fallback son acumulativas y se aplican por POSICION en el
fichero (no por marca de tiempo), que es lo unico que funciona cuando el reloj
vuelve atras y las marcas dejan de ser crecientes. Un CSV que cruce varios
cambios de hora se corrige entero, cada tramo con lo suyo.

El CSV corregido (solo si hizo falta el fallback) se escribe en una copia
temporal (el original no se toca) y es esa copia la que se le pasa a
`convert-meteo -icf`. Los CSV raw son normalmente cuarto-horarios: la correccion
mueve la marca de tiempo una hora entera, o sea 4 filas de paso 15 min, sin tocar
los valores medidos ni el resto del formato.

El resto de minutos va SIEMPRE al `-imt` del convert-meteo (haya hecho falta o no
el fallback), y Fabric (`csv_ingest`) deshace ese `-imt` al ingerir el CSV de
resultados, para que la serie simulada vuelva a la escala de tiempo del SCADA.
Las horas que resuelve `UsesDST` NO hace falta deshacerlas en Fabric: el CSV raw
nunca se tocó para ellas.

Este script NO escribe en el MEF. Si el MEF trae `HourShiftF` o `TimeShiftF`, eso
es un desfase de mas que nadie controla: se avisa por el log y no se toca.

Uso desde auto_run.bat (una llamada por CSV, antes del convert-meteo):

    python timeshift.py aplicar FRA --ws "%WS%" --csv "<csv>" --out "<fichero>"

Por stdout (que el .bat manda al log) sale UNA linea con lo aplicado, mas los
avisos si los hay. El detalle completo se guarda y solo se escribe si algo falla
-- o si se pide con `-v`, para trastear a mano. En `--out` queda un fichero
clave=valor con lo que el .bat necesita:

    imt=6
    csv=C:\...\Temp\pvsyst_tsfix\FRA_RawData_20260421-20260730.csv

Consultas, sin tocar nada:

    python timeshift.py mostrar                 # T08 de todas las plantas
    python timeshift.py mostrar FRA PL2
    python timeshift.py revisar FRA --csv <csv>  # que hay en las fechas de T08

Codigos de salida: 0 si todo bien; 1 si no se pudo determinar el desfase o
escribir el CSV corregido (en ese caso auto_run.bat aborta ESE csv y lo reintenta
en el siguiente ciclo, en vez de simular con un desfase equivocado).

Dependencias: ninguna nueva. Usa openpyxl si esta instalado y, si no, un lector
de .xlsx con la libreria estandar (ver _filas_tabla_stdlib).
"""
from __future__ import annotations

import argparse
import datetime as dt
import os
import re
import sys
import tempfile
from collections import Counter
from pathlib import Path

try:
    import yaml
except ImportError:                                     # pragma: no cover
    yaml = None

# Los mismos defaults que tienen cableados los .bat.
WS_DEFAULT = r"C:\Users\PVSYST1\PVsyst8.1_Data"

# Rango que admite el -imt del CLI. El resto de la particion siempre cae aqui.
IMT_MIN, IMT_MAX = -30, 30

NOMBRE_TABLA = "T08_time_shifts"
NOMBRE_EXCEL = "plants_config.xlsx"

# Carpeta del Excel dentro de la biblioteca SharePoint (Documentos/). Es la misma
# ruta que ve Fabric a traves del shortcut: Files/Raw PVSyst1/CONFIG/.
CARPETA_CONFIG_SP = "CONFIG"

# Donde se deja el CSV corregido. Nombre de fichero IGUAL que el original, para
# que la cabecera del MET y de los resultados siga citando el CSV de verdad.
CARPETA_TSFIX = "pvsyst_tsfix"

PREF = "    "        # sangrado para alinear con las lineas que escribe el .bat


# --------------------------------------------------------------------------
# Config
# --------------------------------------------------------------------------
def cargar_config() -> dict:
    """config.yaml del directorio del script. {} si no hay (o no hay PyYAML)."""
    ruta = Path(__file__).with_name("config.yaml")
    if yaml is None or not ruta.exists():
        return {}
    with open(ruta, "r", encoding="utf-8") as f:
        return yaml.safe_load(f) or {}


def _ruta_cache(cfg: dict) -> Path:
    """Donde se guarda la copia local del Excel de configuracion.

    Sirve de red de seguridad: si Graph falla puntualmente se usa la ultima copia
    buena en vez de abortar la corrida (con un aviso bien visible en el log).
    """
    custom = (cfg.get("paths") or {}).get("plants_config_cache")
    if custom:
        return Path(custom)
    state = (cfg.get("paths") or {}).get("state_file")
    base = Path(state).parent if state else Path(__file__).parent
    return base / "plants_config.cache.xlsx"


def localizar_excel(cfg: dict, forzar_descarga: bool = False) -> tuple[Path, str]:
    """Devuelve (ruta_local_del_excel, de_donde_viene).

    Prioridad:
      1. `paths.plants_config` del config.yaml, si existe en disco (copia local
         que el usuario mantenga a mano).
      2. Descarga via Graph de <CONFIG>/plants_config.xlsx a la cache local.
      3. La cache de una descarga anterior, si la de ahora ha fallado.

    Lanza RuntimeError si no hay ninguna de las tres.
    """
    local = (cfg.get("paths") or {}).get("plants_config")
    if local and not forzar_descarga:
        p = Path(local)
        if p.is_file():
            return p, f"copia local {p}"

    cache = _ruta_cache(cfg)
    carpeta = ((cfg.get("graph") or {}).get("config_folder") or CARPETA_CONFIG_SP).strip("/")
    remoto = f"{carpeta}/{NOMBRE_EXCEL}" if carpeta else NOMBRE_EXCEL

    error_graph = None
    try:
        from sharepoint import descargar_fichero      # import perezoso: solo si toca bajar
        descargar_fichero(cfg, remoto, cache)
        return cache, f"SharePoint {remoto} (via Graph)"
    except Exception as e:
        error_graph = e

    if cache.is_file():
        edad = dt.datetime.fromtimestamp(cache.stat().st_mtime)
        print(f"{PREF}  [AVISO] no se pudo bajar {remoto} de SharePoint ({error_graph}).")
        print(f"{PREF}          se usa la copia en cache de {edad:%Y-%m-%d %H:%M}: {cache}")
        return cache, f"cache local de {edad:%Y-%m-%d %H:%M}"

    raise RuntimeError(
        f"no hay forma de leer {NOMBRE_EXCEL}: Graph fallo ({error_graph}) y no hay "
        f"cache en {cache}. Deja una copia y apuntala con paths.plants_config "
        f"en config.yaml, o arregla el acceso a SharePoint."
    )


# --------------------------------------------------------------------------
# Lectura de T08_time_shifts
# --------------------------------------------------------------------------
class Regla:
    """Una fila de T08.

    `minutos` es el desfase TOTAL del tramo. `inicio` y `fin` son las fechas de
    los cambios de hora que lo abren y lo cierran: en esos dias es donde se busca
    la hora repetida o la hora que falta. None = sin limite por ese lado.
    """

    __slots__ = ("inicio", "fin", "minutos", "fila")

    def __init__(self, inicio, fin, minutos, fila):
        self.inicio = inicio
        self.fin = fin
        self.minutos = minutos
        self.fila = fila

    def cubre(self, dia: dt.date) -> bool:
        if self.inicio is not None and dia < self.inicio:
            return False
        if self.fin is not None and dia > self.fin:
            return False
        return True

    def fechas(self) -> list[dt.date]:
        """Las fechas donde toca mirar el CSV, en orden."""
        return [f for f in (self.inicio, self.fin) if f is not None]

    def __repr__(self):
        i = self.inicio.isoformat() if self.inicio else "(sin inicio)"
        f = self.fin.isoformat() if self.fin else "(sin fin)"
        return f"{i} .. {f} -> {self.minutos:+d} min"


def _a_fecha(valor):
    """Fecha de una celda: datetime, date, numero de serie de Excel o texto.

    Sin openpyxl las fechas llegan como el numero de serie de Excel (dias desde
    el 30/12/1899, con el bug del 1900 bisiesto ya incorporado para serie > 60),
    porque el lector de la libreria estandar no interpreta estilos. Devuelve None
    si la celda esta vacia o no hay forma de leer una fecha.
    """
    if valor is None:
        return None
    if isinstance(valor, bool):
        return None
    if isinstance(valor, dt.datetime):
        return valor.date()
    if isinstance(valor, dt.date):
        return valor
    if isinstance(valor, (int, float)):
        # Por debajo de 61 no se puede distinguir el bug del 1900 bisiesto; nadie
        # pone fechas de enero de 1900 en T08, asi que se descartan.
        if valor < 61:
            return None
        return (dt.datetime(1899, 12, 30) + dt.timedelta(days=float(valor))).date()
    texto = str(valor).strip()
    if not texto:
        return None
    # Solo la parte de fecha: en T08 las horas no aportan nada (los cambios de
    # hora se localizan a dia completo) y vienen de formas muy distintas segun
    # quien edite la celda.
    texto = texto.split("T")[0].split(" ")[0]
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y", "%m/%d/%Y", "%Y/%m/%d", "%d-%m-%Y"):
        try:
            return dt.datetime.strptime(texto, fmt).date()
        except ValueError:
            continue
    return None


def _a_entero(valor):
    """Minutos de una celda. None si esta vacia o no es un numero."""
    if valor is None:
        return None
    if isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float)):
        try:
            if valor != valor:        # NaN
                return None
        except TypeError:
            return None
        return int(round(valor))
    texto = str(valor).strip().replace(",", ".")
    if not texto:
        return None
    try:
        return int(round(float(texto)))
    except ValueError:
        return None


def _filas_tabla(xlsx: Path, nombre_tabla: str) -> list[list]:
    """Filas de una tabla del Excel por su NOMBRE (no por hoja).

    Con openpyxl si esta instalado; si no, con el lector de la libreria estandar
    de mas abajo (un .xlsx es un zip de XML). El PC de PVsyst solo tiene
    garantizados PyYAML y requests, y quedarse sin desfase corta el ciclo diario
    entero, asi que esto no puede depender de un paquete opcional.
    """
    try:
        import openpyxl
    except ImportError:
        return _filas_tabla_stdlib(xlsx, nombre_tabla)

    wb = openpyxl.load_workbook(xlsx, data_only=True)
    try:
        objetivo = nombre_tabla.lower()
        for hoja in wb.worksheets:
            mapa = {str(k).lower(): k for k in hoja.tables}
            if objetivo in mapa:
                tabla = hoja.tables[mapa[objetivo]]
                return [[c.value for c in fila] for fila in hoja[tabla.ref]]

        # Fallback por nombre de hoja.
        for hoja in wb.worksheets:
            t = hoja.title.lower()
            if "time_shift" in t or "time shift" in t or t.startswith(("08", "t08")):
                return [[c.value for c in fila] for fila in hoja.iter_rows()]
    finally:
        wb.close()

    raise RuntimeError(f"no se encontro la tabla '{nombre_tabla}' en {xlsx}")


# --------------------------------------------------------------------------
# Lector de .xlsx con la libreria estandar (sin openpyxl)
# --------------------------------------------------------------------------
# Un .xlsx es un zip de XML. Solo hace falta un trozo muy pequeno del formato:
# localizar la tabla por displayName, su rango, y los valores de esas celdas.
#
# Lo que NO se interpreta: estilos. Asi que las fechas llegan como el numero de
# serie de Excel y las convierte _a_fecha(). Es el unico punto donde este lector
# se aparta de openpyxl, y para T08 da igual porque sabemos por columna cual es
# fecha y cual no.
NS_MAIN = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
NS_REL  = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def _ref_a_fila_col(ref: str) -> tuple[int, int]:
    """'BC12' -> (12, 55). Fila y columna en base 1."""
    letras = "".join(c for c in ref if c.isalpha()).upper()
    digitos = "".join(c for c in ref if c.isdigit())
    col = 0
    for c in letras:
        col = col * 26 + (ord(c) - 64)
    return int(digitos or 0), col


def _rango(ref: str) -> tuple[int, int, int, int]:
    """'A1:D3' -> (fila_ini, col_ini, fila_fin, col_fin)."""
    trozos = ref.replace("$", "").split(":")
    f1, c1 = _ref_a_fila_col(trozos[0])
    f2, c2 = _ref_a_fila_col(trozos[-1])
    return min(f1, f2), min(c1, c2), max(f1, f2), max(c1, c2)


def _rels(z, ruta_xml: str) -> dict:
    """{rId: destino_absoluto_en_el_zip} del fichero .rels de `ruta_xml`."""
    partes = ruta_xml.rsplit("/", 1)
    base = partes[0] if len(partes) == 2 else ""
    rels = f"{base}/_rels/{partes[-1]}.rels" if base else f"_rels/{partes[-1]}.rels"
    if rels not in z.namelist():
        return {}
    import xml.etree.ElementTree as ET
    salida = {}
    for r in ET.fromstring(z.read(rels)):
        destino = r.attrib.get("Target", "")
        if destino.startswith("/"):
            destino = destino.lstrip("/")
        else:
            # Resolver los '../' relativos al directorio de ruta_xml.
            trozos = [t for t in base.split("/") if t] if base else []
            for pieza in destino.split("/"):
                if pieza == "..":
                    if trozos:
                        trozos.pop()
                elif pieza not in ("", "."):
                    trozos.append(pieza)
            destino = "/".join(trozos)
        salida[r.attrib.get("Id", "")] = destino
    return salida


def _cadenas_compartidas(z) -> list[str]:
    """xl/sharedStrings.xml -> lista de textos (las celdas de texto lo indexan)."""
    if "xl/sharedStrings.xml" not in z.namelist():
        return []
    import xml.etree.ElementTree as ET
    salida = []
    for si in ET.fromstring(z.read("xl/sharedStrings.xml")):
        # El texto puede venir partido en varios <r><t> (texto enriquecido).
        salida.append("".join(t.text or "" for t in si.iter(f"{NS_MAIN}t")))
    return salida


def _valor_celda(celda, cadenas):
    """Valor de un <c>, sin mirar estilos: texto, numero, bool o None."""
    tipo = celda.attrib.get("t", "n")
    if tipo == "inlineStr":
        return "".join(t.text or "" for t in celda.iter(f"{NS_MAIN}t")) or None
    v = celda.find(f"{NS_MAIN}v")
    if v is None or v.text is None:
        return None
    texto = v.text
    if tipo == "s":
        try:
            return cadenas[int(texto)]
        except (ValueError, IndexError):
            return None
    if tipo == "b":
        return texto not in ("0", "", "false", "FALSE")
    if tipo == "e":
        return None                       # celda en error (#REF!, #N/A...)
    if tipo in ("str", "d"):
        return texto
    try:
        return float(texto)
    except ValueError:
        return texto


def _filas_de_hoja(z, ruta_hoja: str, cadenas, ref=None) -> list[list]:
    """Filas de una hoja (o de un rango de ella) como lista de listas."""
    import xml.etree.ElementTree as ET
    if ref:
        f_ini, c_ini, f_fin, c_fin = _rango(ref)
    else:
        f_ini, c_ini, f_fin, c_fin = 1, 1, 10 ** 9, 0

    celdas: dict[int, dict[int, object]] = {}
    raiz = ET.fromstring(z.read(ruta_hoja))
    for fila in raiz.iter(f"{NS_MAIN}row"):
        for celda in fila.iter(f"{NS_MAIN}c"):
            f, c = _ref_a_fila_col(celda.attrib.get("r", ""))
            if not f or f < f_ini or f > f_fin:
                continue
            if ref and (c < c_ini or c > c_fin):
                continue
            valor = _valor_celda(celda, cadenas)
            if valor is not None or ref:
                celdas.setdefault(f, {})[c] = valor
                if not ref:
                    c_fin = max(c_fin, c)

    if not celdas:
        return []
    salida = []
    for f in range(f_ini, max(celdas) + 1):
        fila_celdas = celdas.get(f, {})
        salida.append([fila_celdas.get(c) for c in range(c_ini, c_fin + 1)])
    return salida


def _filas_tabla_stdlib(xlsx: Path, nombre_tabla: str) -> list[list]:
    """Igual que _filas_tabla() pero solo con la libreria estandar."""
    import xml.etree.ElementTree as ET
    import zipfile

    objetivo = nombre_tabla.lower()
    with zipfile.ZipFile(xlsx) as z:
        libro = ET.fromstring(z.read("xl/workbook.xml"))
        propiedades = libro.find(f"{NS_MAIN}workbookPr")
        if propiedades is not None and propiedades.attrib.get("date1904", "0") in ("1", "true"):
            raise RuntimeError(
                f"{xlsx.name} usa el sistema de fechas 1904 (Excel de Mac), que este "
                f"lector no convierte. Instala openpyxl (`pip install openpyxl`) o "
                f"vuelve a guardar el Excel con el sistema de fechas 1900."
            )

        rels_libro = _rels(z, "xl/workbook.xml")
        cadenas = _cadenas_compartidas(z)

        hojas = []      # [(nombre, ruta_en_el_zip)]
        for h in libro.iter(f"{NS_MAIN}sheet"):
            ruta = rels_libro.get(h.attrib.get(f"{NS_REL}id", ""))
            if ruta and ruta in z.namelist():
                hojas.append((h.attrib.get("name", ""), ruta))

        # 1. Por nombre de tabla (displayName), que es como la busca Fabric.
        for _, ruta in hojas:
            rels_hoja = _rels(z, ruta)
            hoja = ET.fromstring(z.read(ruta))
            for parte in hoja.iter(f"{NS_MAIN}tablePart"):
                ruta_tabla = rels_hoja.get(parte.attrib.get(f"{NS_REL}id", ""))
                if not ruta_tabla or ruta_tabla not in z.namelist():
                    continue
                tabla = ET.fromstring(z.read(ruta_tabla))
                nombres = {(tabla.attrib.get("displayName") or "").lower(),
                           (tabla.attrib.get("name") or "").lower()}
                if objetivo in nombres:
                    return _filas_de_hoja(z, ruta, cadenas, tabla.attrib.get("ref"))

        # 2. Fallback por nombre de hoja.
        for nombre, ruta in hojas:
            t = nombre.lower()
            if "time_shift" in t or "time shift" in t or t.startswith(("08", "t08")):
                return _filas_de_hoja(z, ruta, cadenas)

    raise RuntimeError(f"no se encontro la tabla '{nombre_tabla}' en {xlsx}")


def leer_reglas(xlsx: Path) -> dict[str, list[Regla]]:
    """{PLANTA: [Regla, ...]} desde T08_time_shifts, EN EL ORDEN DE LA TABLA.

    Misma lectura que hace `csv_ingest` en Fabric, para que las dos puntas de la
    cadena no puedan interpretar la tabla de forma distinta:
      - `end_date` es INCLUSIVO a dia completo.
      - Puede haber VARIAS filas por planta (varios cambios de hora).
      - Si dos tramos se solapan gana el de MAS ABAJO en la tabla.
      - Un 0 explicito es una regla valida: anula un tramo anterior solapado.
      - Fecha vacia = sin limite por ese lado (y sin fecha donde buscar el
        cambio de hora, o sea que ese lado no corrige ninguna hora).
      - Fila sin `time_shift` numerico = no es una regla, se ignora.
    """
    filas = _filas_tabla(xlsx, NOMBRE_TABLA)
    if not filas:
        return {}

    cabecera = [str(c).strip().lower() if c is not None else "" for c in filas[0]]

    def indice(*nombres):
        for n in nombres:
            if n in cabecera:
                return cabecera.index(n)
        return None

    i_planta = indice("plant_code", "planta", "plant")
    i_min    = indice("time_shift", "time_shift_min", "minutos", "min")
    i_ini    = indice("start_date", "inicio", "start")
    i_fin    = indice("end_date", "fin", "end")
    if i_planta is None or i_min is None:
        raise RuntimeError(
            f"{NOMBRE_TABLA} necesita al menos las columnas 'plant_code' y "
            f"'time_shift'; se han leido: {', '.join(c for c in cabecera if c)}"
        )

    reglas: dict[str, list[Regla]] = {}
    for n, fila in enumerate(filas[1:], start=2):
        def celda(i):
            return fila[i] if i is not None and i < len(fila) else None

        planta = celda(i_planta)
        if planta is None or not str(planta).strip():
            continue
        minutos = _a_entero(celda(i_min))
        if minutos is None:
            continue
        reglas.setdefault(str(planta).strip().upper(), []).append(
            Regla(_a_fecha(celda(i_ini)), _a_fecha(celda(i_fin)), minutos, n)
        )
    return reglas


def shift_en(reglas: list[Regla], dia: dt.date) -> int:
    """Minutos en vigor ese dia. Gana la ULTIMA regla de la tabla que lo cubre."""
    total = 0
    for r in reglas or ():
        if r.cubre(dia):
            total = r.minutos
    return total


# --------------------------------------------------------------------------
# Particion horas / minutos
# --------------------------------------------------------------------------
def descomponer(total: int) -> tuple[int, int]:
    """total -> (horas, resto) con |resto| <= 30 y horas*60 + resto == total.

    En los empates (30, 90, 150...) se prefiere el reparto con MENOS horas, asi
    que el resto se queda en +-30, que el CLI todavia acepta:
        90 -> (1, 30)      -90 -> (-1, -30)      68 -> (1, 8)
    """
    signo = -1 if total < 0 else 1
    horas = signo * ((abs(total) + IMT_MAX - 1) // 60)
    resto = total - horas * 60
    if not IMT_MIN <= resto <= IMT_MAX:
        # No deberia pasar nunca; si pasara, el CLI recibiria un -imt fuera de
        # rango y el MET saldria desalineado sin protestar. Mejor romper aqui.
        raise ValueError(f"reparto invalido de {total} min: {horas} h y {resto} min")
    return horas, resto


# --------------------------------------------------------------------------
# Formato del CSV raw, declarado en el MEF
# --------------------------------------------------------------------------
# El MEF se LEE (nunca se escribe) porque es lo que usa de verdad el
# convert-meteo: separador, linea donde empiezan los datos, formato de la fecha y
# en que columna esta. Los campos van por POSICION, no por nombre.
def leer_formato_mef(mef: Path) -> dict:
    """Formato del CSV raw declarado en el MEF: separador, cabecera, paso, campos."""
    texto = mef.read_text(encoding="cp1252", errors="replace")

    def campo(clave, defecto=None):
        m = re.search(r"^\s*" + re.escape(clave) + r"=(.*)$", texto, re.MULTILINE)
        return m.group(1).strip() if m else defecto

    separ = campo("Separ", "$002C")
    if separ.startswith("$"):
        try:
            separ = chr(int(separ[1:], 16))
        except ValueError:
            separ = ","
    elif separ.lower() in ("tab", "tabulador"):
        separ = "\t"
    else:
        separ = (separ or ",")[0]

    # NoFieldDate=1,0,0,0,0, -> el primer numero es la columna de la fecha.
    campo_fecha = 1
    nfd = campo("NoFieldDate")
    if nfd:
        try:
            campo_fecha = int(nfd.split(",")[0])
        except (ValueError, IndexError):
            pass

    # Bloques VarASCII: VarIdent=GlobHor ... NoField=2
    variables = {}
    for bloque in re.findall(r"VarASCII=\d+(.*?)End of TVarASCII", texto, re.DOTALL):
        ident = re.search(r"VarIdent=(\S+)", bloque)
        nof   = re.search(r"NoField=(\d+)", bloque)
        mult  = re.search(r"FieldMult=([\d.eE+-]+)", bloque)
        if ident and nof:
            try:
                factor = float(mult.group(1)) if mult else 1.0
            except ValueError:
                factor = 1.0
            variables[ident.group(1)] = (int(nof.group(1)), factor)

    return {
        "separador": separ,
        "cabecera": _a_entero(campo("NLigDeb", "1")) or 1,
        "paso": _a_entero(campo("MeasureStep", "15")) or 15,
        "campo_fecha": campo_fecha,
        "orden_fecha": campo("DateFmtType", "DDxMMxYYxhhxmm"),
        "variables": variables,
        "hourshiftf": campo("HourShiftF"),
        "timeshiftf": campo("TimeShiftF"),
        # Campos de los que depende es_dst_fiable(): si PVsyst ya resuelve el
        # cambio de hora solo, no hace falta (ni conviene) reescribir el CSV.
        "usesdst": campo("UsesDST"),
        "reftimes": campo("RefTimeS"),
        "datehete": _a_fecha(campo("DateHEte")),
        "datehhiver": _a_fecha(campo("DateHHiver")),
    }


def _partes_fecha(cadena: str) -> list[int]:
    """Grupos de digitos de una marca de tiempo, en orden de aparicion."""
    return [int(g) for g in re.findall(r"\d+", cadena)]


def parsear_fecha(cadena: str, orden: str):
    """Marca de tiempo segun el DateFmtType del MEF (p.ej. MMxDDxYYxhhxmm).

    Se trocea en grupos de digitos y se reparten segun el orden declarado, asi
    da igual que el separador sea '/', '-' o ':'. Devuelve None si no cuadra.
    """
    partes = _partes_fecha(cadena)
    claves = re.findall(r"YYYY|YY|MM|DD|hh|mm|ss", orden)
    if len(partes) < 5 or len(claves) < 5:
        return None
    campos = {}
    for clave, valor in zip(claves, partes):
        campos[clave[:2] if clave != "YYYY" else "YY"] = valor
    try:
        anio = campos["YY"]
        if anio < 100:
            anio += 2000
        hora, minuto = campos.get("hh", 0), campos.get("mm", 0)
        base = dt.datetime(anio, campos["MM"], campos["DD"])
        # PVsyst escribe 24:00 para el final del dia.
        return base + dt.timedelta(hours=hora, minutes=minuto, seconds=campos.get("ss", 0))
    except (KeyError, ValueError):
        return None


def rehacer_fecha(texto: str, nuevo: dt.datetime, orden: str) -> str:
    """Escribe `nuevo` con la MISMA forma textual que `texto`.

    Se sustituye grupo de digitos por grupo de digitos, conservando separadores y
    relleno con ceros: '4/21/2026 0:15' con una hora mas sale '4/21/2026 1:15', y
    '04/21/2026 00:15' sale '04/21/2026 01:15'. Reformatear a un patron canonico
    seria pedirle al convert-meteo que se tragase un formato que no es el que
    declara el MEF.
    """
    claves = re.findall(r"YYYY|YY|MM|DD|hh|mm|ss", orden)
    valores = {"MM": nuevo.month, "DD": nuevo.day, "hh": nuevo.hour,
               "mm": nuevo.minute, "ss": nuevo.second}
    salida, i = [], 0
    for trozo in re.split(r"(\d+)", texto):
        if trozo.isdigit() and i < len(claves):
            clave = claves[i]
            if clave in ("YY", "YYYY"):
                # El ancho del original manda: 2 digitos -> año corto.
                valor = nuevo.year if len(trozo) > 2 else nuevo.year % 100
            else:
                valor = valores.get(clave, int(trozo))
            salida.append(str(valor).zfill(len(trozo)))
            i += 1
        else:
            salida.append(trozo)
    return "".join(salida)


def _leer_lineas(csv: Path) -> list[str]:
    """Lineas del CSV con su terminador intacto (latin-1: ida y vuelta sin perder
    ningun byte, y aqui solo se toca la marca de tiempo, que es ASCII)."""
    return csv.read_bytes().decode("latin-1").splitlines(keepends=True)


class Marca:
    """Una marca de tiempo del CSV: en que linea esta, su texto y su valor."""

    __slots__ = ("linea", "texto", "valor")

    def __init__(self, linea, texto, valor):
        self.linea = linea          # indice en la lista de lineas (0-based)
        self.texto = texto          # el texto tal cual esta en el fichero
        self.valor = valor          # datetime


def leer_marcas(lineas: list[str], fmt: dict) -> tuple[list[Marca], int]:
    """Marcas de tiempo del CSV, en orden de fichero. Devuelve (marcas, ilegibles)."""
    sep, col = fmt["separador"], fmt["campo_fecha"]
    marcas, malas = [], 0
    for i, linea in enumerate(lineas):
        if i < fmt["cabecera"]:
            continue
        if not linea.strip():
            continue
        trozos = linea.rstrip("\r\n").split(sep)
        if len(trozos) < col:
            malas += 1
            continue
        texto = trozos[col - 1]
        valor = parsear_fecha(texto, fmt["orden_fecha"])
        if valor is None:
            malas += 1
            continue
        marcas.append(Marca(i, texto, valor))
    return marcas, malas


# --------------------------------------------------------------------------
# Deteccion del cambio de hora en una fecha concreta
# --------------------------------------------------------------------------
# Un CSV real trae muchos huecos que NO son cambios de hora (cortes del SCADA de
# horas o dias). Por eso solo se mira en las fechas que dice T08, y solo se acepta
# como cambio de hora un hueco de tamano compatible con una hora.
HUECO_MAX = 120         # minutos: por encima de esto es un corte de datos
HUECO_OBJETIVO = 60     # minutos: lo que mide el salto de un cambio de hora

# Franja horaria donde puede caer un cambio de hora. En la UE es a las 02:00/03:00
# locales, asi que con [00:00, 06:00) sobra de margen. Sin esta ventana, un corte
# del SCADA de exactamente una hora a media tarde se confundiria con un cambio de
# hora: el CSV de PL2 tiene uno el 17/05 a las 18:45, por ejemplo.
VENTANA_DST = (0, 6)

# Hora del cambio (UE). Sirve para situar el instante exacto del cambio y decidir
# si cae dentro del CSV, antes o despues. csv_ingest usa la misma constante para
# cortar los tramos al revertir.
HORA_CAMBIO = 2


class Artefacto:
    """Lo que se ha encontrado en la fecha que senala T08."""

    __slots__ = ("tipo", "linea", "instante", "detalle", "horas")

    def __init__(self, tipo, linea, instante, detalle, horas):
        self.tipo = tipo            # 'hueco' | 'repetida' | 'previo'
        self.linea = linea          # linea desde la que se corrige (incluida)
        self.instante = instante
        self.detalle = detalle
        self.horas = horas          # horas a sumar (negativo = restar)


def eventos_de_cambio(reglas: list[Regla]) -> list[tuple]:
    """[(fecha, papel, horas)] de los cambios de hora que declara T08.

    `papel` dice de que lado del tramo esta la fecha, y con ello en que sentido
    movio el reloj:
      - 'inicio': lo abre, el reloj se ADELANTA  -> falta una hora en el CSV
      - 'fin':    lo cierra, el reloj vuelve ATRAS -> hora repetida en el CSV
    `horas` son las horas enteras del tramo (de descomponer(minutos)). Un tramo sin
    horas enteras no tiene cambio de hora que corregir y no genera eventos.
    """
    eventos = set()
    for r in reglas:
        horas, _ = descomponer(r.minutos)
        if horas == 0:
            continue
        if r.inicio is not None:
            eventos.add((r.inicio, "inicio", horas))
        if r.fin is not None:
            eventos.add((r.fin, "fin", horas))
    return sorted(eventos, key=lambda e: (e[0], e[1]))


def instante_de_cambio(fecha: dt.date) -> dt.datetime:
    """El momento exacto en que el reloj cambia ese dia."""
    return dt.datetime.combine(fecha, dt.time(HORA_CAMBIO))


# Sentido en que movio el reloj cada lado del tramo, y el artefacto que deja.
SIGNO_PAPEL = {"inicio": -1, "fin": +1}
TIPO_ESPERADO = {"inicio": "hueco", "fin": "repetida"}


# --------------------------------------------------------------------------
# Deteccion de si UsesDST es de fiar (si no, se cae al fallback de reescribir
# el CSV a mano)
# --------------------------------------------------------------------------
def _ultimo_domingo(anio: int, mes: int) -> dt.date:
    """Ultimo domingo de `mes`/`anio` (la regla de la UE para el cambio de hora)."""
    if mes == 12:
        siguiente = dt.date(anio + 1, 1, 1)
    else:
        siguiente = dt.date(anio, mes + 1, 1)
    ultimo_dia = siguiente - dt.timedelta(days=1)
    return ultimo_dia - dt.timedelta(days=(ultimo_dia.weekday() - 6) % 7)


def es_dst_fiable(fmt: dict, fecha_evento: dt.date) -> tuple[bool, str]:
    """Si PVsyst va a resolver el cambio de hora de `fecha_evento` el solo.

    Comprobado con pruebas reales (FRA, PL1, PL2, 31/07/2026): cuando la MEF
    declara `UsesDST=True` + `RefTimeS=LegalTime` + `DateHEte`/`DateHHiver`
    coherentes con la regla de la UE (ultimo domingo de marzo/octubre, del año
    del propio cambio), `run-simulation` ya corrige ese cambio de hora sin que
    haga falta tocar el CSV -- y tocarlo ademas DUPLICA la correccion.

    Si algo de esto no cuadra, no hay garantia de que PVsyst vaya a hacer nada:
    se devuelve False y el llamador cae al mecanismo antiguo (buscar el hueco o
    la hora repetida en el CSV y reescribirlo a mano).

    Devuelve (fiable, motivo) -- el motivo se registra siempre en el log, tanto
    si es de fiar como si no, para poder diagnosticar sin repetir la corrida.
    """
    usesdst = str(fmt.get("usesdst") or "").strip().lower()
    if usesdst != "true":
        return False, f"UsesDST no es True en la MEF (es '{fmt.get('usesdst')}')"

    reftimes = str(fmt.get("reftimes") or "").strip().lower()
    if reftimes != "legaltime":
        return False, f"RefTimeS no es LegalTime en la MEF (es '{fmt.get('reftimes')}')"

    dhe, dhh = fmt.get("datehete"), fmt.get("datehhiver")
    if dhe is None or dhh is None:
        return False, "faltan DateHEte/DateHHiver (o no se han podido leer) en la MEF"
    if dhe >= dhh:
        return False, f"DateHEte ({dhe.isoformat()}) no es anterior a DateHHiver ({dhh.isoformat()})"

    anio = fecha_evento.year
    if dhe.year != anio and dhh.year != anio:
        return False, (f"DateHEte/DateHHiver son de {dhe.year}, no del año del cambio "
                        f"({anio}): revisa que se hayan actualizado este año en la MEF")

    esperada_marzo = _ultimo_domingo(anio, 3)
    esperada_octubre = _ultimo_domingo(anio, 10)
    if dhe != esperada_marzo or dhh != esperada_octubre:
        return False, (f"DateHEte/DateHHiver ({dhe.isoformat()}/{dhh.isoformat()}) no "
                        f"coinciden con la regla de la UE para {anio} "
                        f"({esperada_marzo.isoformat()}/{esperada_octubre.isoformat()})")

    return True, (f"UsesDST=True, RefTimeS=LegalTime, DateHEte/DateHHiver "
                   f"({dhe.isoformat()}/{dhh.isoformat()}) coherentes con la regla de la UE")


def cambio_anterior_al_csv(marcas: list[Marca], fecha: dt.date, papel: str,
                           horas: int) -> Artefacto:
    """Cambio de hora ANTERIOR al CSV: su efecto ya esta en todas las filas.

    Si el reloj cambio antes de que empiecen los datos, en el fichero no hay ni
    hora repetida ni hora que falte -- pero todas las marcas vienen ya con el
    desfase puesto. Asi que la correccion se aplica desde la PRIMERA fila, y el
    sentido lo da el lado del tramo (no se puede sacar de un artefacto que no
    existe).

    Con las dos fechas del tramo por detras del CSV las dos correcciones se suman
    y se anulan, que es lo correcto: el tramo ya se abrio Y se cerro antes de que
    empiecen los datos.
    """
    signo = SIGNO_PAPEL[papel]
    return Artefacto("previo", marcas[0].linea, marcas[0].valor,
                     f"el cambio del {fecha.isoformat()} ({papel} del tramo) es anterior "
                     f"al primer dato del CSV ({marcas[0].valor:%Y-%m-%d %H:%M}), asi que "
                     f"su desfase ya viene aplicado en todo el fichero",
                     signo * horas)


def detectar_cambio_hora(marcas: list[Marca], dia: dt.date, paso: int, horas: int = 1):
    """Busca en `dia` una hora repetida o una hora que falte.

    Devuelve (Artefacto | None, motivo_texto). El motivo se registra siempre en el
    log, tanto si hay artefacto como si no: es la traza de por que se ha corregido
    (o de por que no).

    El SENTIDO lo manda lo que hay en el CSV, no el signo de T08:
      - falta una hora (el reloj salta adelante)  -> se RESTAN `horas` despues
      - hora repetida  (el reloj vuelve atras)    -> se SUMAN `horas` despues
    En los dos casos la serie queda continua, en hora estandar. La MAGNITUD sale de
    T08 (`horas`), para que el desplazamiento y el revert de Fabric coincidan.
    """
    indices = [i for i, m in enumerate(marcas) if m.valor.date() == dia]
    if not indices:
        return None, f"el CSV no tiene datos del {dia.isoformat()}"

    def de_madrugada(marca) -> bool:
        return VENTANA_DST[0] <= marca.valor.hour < VENTANA_DST[1]

    # --- hora repetida: el reloj vuelve atras ---
    # Primero por marcas no crecientes (es lo que deja un reloj que retrocede) y,
    # si el SCADA las ha reordenado, por etiquetas duplicadas.
    for i in indices:
        if not de_madrugada(marcas[i]):
            continue
        if i > 0 and marcas[i].valor <= marcas[i - 1].valor:
            retroceso = int((marcas[i - 1].valor - marcas[i].valor).total_seconds() // 60)
            return (Artefacto("repetida", marcas[i].linea, marcas[i].valor,
                              f"la marca {marcas[i].valor:%H:%M} viene despues de "
                              f"{marcas[i-1].valor:%H:%M} (retroceso de {retroceso} min)",
                              +horas),
                    None)

    repetidas = [t for t, n in Counter(marcas[i].valor for i in indices).items()
                 if n > 1 and VENTANA_DST[0] <= t.hour < VENTANA_DST[1]]
    if repetidas:
        primera = min(repetidas)
        # Se corrige desde la SEGUNDA aparicion de la primera marca repetida.
        vistas = 0
        for i in indices:
            if marcas[i].valor == primera:
                vistas += 1
                if vistas == 2:
                    return (Artefacto("repetida", marcas[i].linea, primera,
                                      f"{len(repetidas)} marca(s) repetida(s), "
                                      f"la primera {primera:%H:%M}", +horas),
                            None)

    # --- hora que falta: el reloj salta adelante ---
    candidatos, fuera_de_ventana = [], []
    for a, b in zip(indices, indices[1:]):
        salto = int((marcas[b].valor - marcas[a].valor).total_seconds() // 60)
        if not paso < salto <= HUECO_MAX:
            continue
        if de_madrugada(marcas[b]):
            candidatos.append((abs(salto - HUECO_OBJETIVO), salto, a, b))
        else:
            fuera_de_ventana.append(f"{marcas[a].valor:%H:%M}->{marcas[b].valor:%H:%M} ({salto} min)")
    if candidatos:
        _, salto, a, b = min(candidatos)
        return (Artefacto("hueco", marcas[b].linea, marcas[b].valor,
                          f"salto de {marcas[a].valor:%H:%M} a {marcas[b].valor:%H:%M} "
                          f"({salto} min, se esperaban {paso})", -horas),
                None)

    n = len(indices)
    esperadas = (24 * 60) // paso if paso else 0
    grandes = [1 for a, b in zip(indices, indices[1:])
               if (marcas[b].valor - marcas[a].valor).total_seconds() // 60 > HUECO_MAX]
    extra = ""
    if fuera_de_ventana:
        extra += (f"; hay hueco(s) de menos de {HUECO_MAX} min fuera de la franja "
                  f"{VENTANA_DST[0]:02d}:00-{VENTANA_DST[1]:02d}:00 "
                  f"({', '.join(fuera_de_ventana[:3])}): corte de datos, no cambio de hora")
    if grandes:
        extra += (f"; {len(grandes)} hueco(s) de mas de {HUECO_MAX} min "
                  f"(corte de datos, no cambio de hora)")
    return None, (f"el {dia.isoformat()} no tiene ni hora repetida ni hora que falte "
                  f"({n} marcas, se esperaban {esperadas}){extra}")


def _partir_linea(linea: str) -> tuple[str, str]:
    """(cuerpo, terminador) de una linea, para poder rehacerla igual."""
    for term in ("\r\n", "\n", "\r"):
        if linea.endswith(term):
            return linea[:-len(term)], term
    return linea, ""


def corregir_csv(csv: Path, lineas: list[str], fmt: dict,
                 correcciones: list[Artefacto], destino: Path) -> tuple[int, list[str]]:
    """Reescribe el CSV aplicando las correcciones acumuladas. Devuelve (n_filas, log).

    Las correcciones se aplican por POSICION en el fichero: cada una vale desde su
    linea (incluida) hasta el final, y se suman entre ellas. Es lo unico que
    funciona cuando el reloj vuelve atras y las marcas dejan de ser crecientes.

    En la costura de cada corte puede quedar una marca DUPLICADA. Pasa al cerrar un
    hueco: el CSV de PL2 salta de 02:00 a 03:00 el 29/03, asi que al restar la hora
    ese 03:00 se convierte en otro 02:00 y hay dos filas con la misma etiqueta (son
    la misma medida, emitida una vez con la hora de antes del salto y otra con la
    de despues). Esas filas de la costura se eliminan: un meteo con etiquetas
    repetidas es justo lo que hace que PVsyst rechace el fichero o reparta mal los
    intervalos. Se registra en el log cuantas y cuales.
    """
    if not correcciones:
        return 0, []

    ordenadas = sorted(correcciones, key=lambda c: c.linea)
    # delta acumulado en minutos a partir de cada linea de corte
    cortes, acumulado = [], 0
    for c in ordenadas:
        acumulado += c.horas * 60
        cortes.append((c.linea, acumulado))

    def delta(i):
        d = 0
        for linea, acc in cortes:
            if i >= linea:
                d = acc
        return d

    sep, col = fmt["separador"], fmt["campo_fecha"]
    lineas_corte = {c.linea for c in ordenadas}

    nuevas: list[str] = []
    cambiadas = 0
    ultimo = None                 # ultima marca ESCRITA, para detectar la costura
    en_costura = False
    quitadas: list[str] = []

    for i, linea in enumerate(lineas):
        if i < fmt["cabecera"] or not linea.strip():
            nuevas.append(linea)
            continue

        cuerpo, fin = _partir_linea(linea)
        trozos = cuerpo.split(sep)
        valor = parsear_fecha(trozos[col - 1], fmt["orden_fecha"]) if len(trozos) >= col else None
        if valor is None:
            nuevas.append(linea)          # linea que no se entiende: intacta
            continue

        d = delta(i)
        nuevo = valor + dt.timedelta(minutes=d)

        # Al entrar en un corte se abre la ventana de costura, y se cierra en la
        # primera fila que ya no pisa a la anterior.
        if i in lineas_corte:
            en_costura = True
        if en_costura:
            if ultimo is not None and nuevo <= ultimo:
                quitadas.append(f"{valor:%Y-%m-%d %H:%M} (linea {i + 1}) -> {nuevo:%H:%M} duplicado")
                continue
            en_costura = False

        if d:
            trozos[col - 1] = rehacer_fecha(trozos[col - 1], nuevo, fmt["orden_fecha"])
            nuevas.append(sep.join(trozos) + fin)
            cambiadas += 1
        else:
            nuevas.append(linea)
        ultimo = nuevo

    destino.parent.mkdir(parents=True, exist_ok=True)
    tmp = None
    try:
        fd, tmp = tempfile.mkstemp(dir=str(destino.parent), suffix=".tmp")
        with os.fdopen(fd, "wb") as f:
            f.write("".join(nuevas).encode("latin-1"))
        os.replace(tmp, destino)
        tmp = None
    finally:
        if tmp and os.path.exists(tmp):
            try:
                os.remove(tmp)
            except OSError:
                pass

    registro = [f"{c.horas:+d} h desde la linea {c.linea + 1} "
                f"({c.instante:%Y-%m-%d %H:%M}, {c.tipo})" for c in ordenadas]
    if quitadas:
        registro.append(f"{len(quitadas)} fila(s) de la costura eliminada(s) por etiqueta "
                        f"duplicada: {'; '.join(quitadas)}")
    return cambiadas, registro


# --------------------------------------------------------------------------
# Fechas de referencia
# --------------------------------------------------------------------------
RE_FECHA8 = re.compile(r"(?<!\d)(20\d{6})(?!\d)")


def rango_de_nombre(nombre: str) -> tuple[dt.date | None, dt.date | None]:
    """Rango de fechas que anuncia el nombre del CSV de meteo.

    `meteo_export` los nombra `<PLANTA>_RawData_<YYYYMMDD>-<YYYYMMDD>.csv`.
    Con una sola fecha, el rango es ese dia. Sin ninguna, (None, None).
    """
    fechas = []
    for m in RE_FECHA8.finditer(Path(nombre).stem):
        try:
            fechas.append(dt.datetime.strptime(m.group(1), "%Y%m%d").date())
        except ValueError:
            continue
    if not fechas:
        return None, None
    return min(fechas), max(fechas)


# --------------------------------------------------------------------------
# Comandos
# --------------------------------------------------------------------------
def _resto_imt(reglas: list[Regla], ultimo: dt.date) -> tuple[int, list[str]]:
    """Minutos para el -imt: el resto del tramo vigente el ultimo dia del CSV.

    El -imt es UNO por conversion, asi que no puede variar por tramos. Lo normal
    es que el resto sea el mismo en todas las filas de la planta (es el desfase
    propio del SCADA) y solo cambien las horas, que ya se arreglan en el CSV.
    """
    avisos = []
    total = shift_en(reglas, ultimo)
    _, resto = descomponer(total)

    restos = {descomponer(r.minutos)[1] for r in reglas}
    if len(restos) > 1:
        avisos.append(
            f"las filas de T08 no coinciden en los minutos sueltos ({sorted(restos)}): "
            f"el -imt es uno solo para todo el fichero y se usa {resto:+d} "
            f"(el del tramo vigente el {ultimo.isoformat()})")
    if reglas and not any(r.cubre(ultimo) for r in reglas):
        avisos.append(
            f"ninguna fila de T08 cubre el {ultimo.isoformat()} (ultimo dia del CSV), "
            f"asi que el -imt sale 0. Si la planta tiene un desfase propio de minutos, "
            f"anade una fila que cubra ese tramo.")
    return resto, avisos


def _resumen_correccion(art: Artefacto) -> str:
    """Como se describe una correccion en la linea de resumen."""
    if art.tipo == "previo":
        return f"{art.horas:+d} h en todo el fichero"
    return f"{art.horas:+d} h desde {art.instante:%d/%m}"


def cmd_aplicar(args) -> int:
    """Corrige el CSV y devuelve el -imt.

    La salida es UNA linea, porque esto corre por cada CSV de cada ciclo diario y
    el detalle no aporta nada cuando todo va bien. El detalle se guarda igualmente
    y se vuelca entero si algo falla (si no, un fallo seria indiagnosticable) o si
    se pide con --verbose. Los avisos se ven siempre: no son detalle rutinario,
    son cosas que hay que mirar.
    """
    traza: list[str] = []
    avisos: list[str] = []

    def anotar(msg):
        """Detalle: solo se ve con --verbose o si la cosa acaba en error."""
        traza.append(msg)
        if args.verbose:
            print(f"{PREF}  {msg}")

    def avisar(msg):
        traza.append(f"[AVISO] {msg}")
        avisos.append(msg)

    def fallar(msg):
        """Error: se vuelca el detalle acumulado para poder diagnosticarlo."""
        print(f"{PREF}[TIME-SHIFT] {args.planta.upper()}: [ERROR] {msg}")
        for linea in traza:
            print(f"{PREF}  {linea}")
        return 1

    cfg = cargar_config()
    planta = args.planta.upper()
    ws = Path(args.ws or (cfg.get("pvsyst") or {}).get("workspace") or WS_DEFAULT)
    mef = Path(args.mef) if args.mef else ws / "Meteo" / f"{planta}.MEF"
    csv = Path(args.csv) if args.csv else None

    try:
        xlsx, procedencia = localizar_excel(cfg)
        reglas = leer_reglas(xlsx).get(planta, [])
    except Exception as e:
        return fallar(str(e))

    anotar(f"{NOMBRE_TABLA}: {procedencia}")
    for r in reglas:
        anotar(f"  fila {r.fila}: {r}")
    if not reglas:
        anotar(f"  {planta} no tiene reglas -> sin desfase")

    if csv is None or not csv.is_file():
        return fallar(f"hace falta el CSV raw (--csv) y no existe: {csv}")
    if not mef.is_file():
        return fallar(f"no existe el MEF {mef}")

    # --- formato del CSV, segun el MEF ---
    try:
        fmt = leer_formato_mef(mef)
        lineas = _leer_lineas(csv)
        marcas, ilegibles = leer_marcas(lineas, fmt)
    except Exception as e:
        return fallar(f"no se pudo leer el CSV con el formato del MEF: {e}")
    if not marcas:
        return fallar(f"no se ha podido leer ninguna marca de tiempo de {csv.name} "
                      f"(separador '{fmt['separador']}', columna {fmt['campo_fecha']}, "
                      f"formato {fmt['orden_fecha']})")
    anotar(f"CSV: {csv.name}  {len(marcas):,} filas "
           f"{marcas[0].valor:%Y-%m-%d %H:%M} .. {marcas[-1].valor:%Y-%m-%d %H:%M}"
           + (f"  ({ilegibles} lineas ilegibles)" if ilegibles else ""))

    # El MEF ya no se toca, pero si trae desfase propio suma por detras de todo
    # esto y desalinea el MET sin que se vea en ningun sitio.
    for clave, etiqueta in (("hourshiftf", "HourShiftF"), ("timeshiftf", "TimeShiftF")):
        valor = fmt.get(clave)
        if valor and _a_entero(valor) not in (0, None):
            avisar(f"el MEF declara {etiqueta}={valor}: es un desfase de mas que este "
                   f"script NO gestiona. Quitalo del MEF salvo que lo quieras a proposito.")

    # --- correccion de las horas enteras en las fechas de T08 ---
    # Cada fecha de T08 es un cambio de hora, y puede caer en tres sitios:
    #   dentro del CSV  -> se busca la hora repetida / la que falta
    #   antes del CSV   -> su desfase ya viene aplicado: se corrige desde la 1a fila
    #   despues del CSV -> ese cambio no le afecta a este fichero
    eventos = eventos_de_cambio(reglas)
    correcciones: list[Artefacto] = []
    usesdst_resuelve: list[str] = []       # fechas que deja resolver a PVsyst
    hubo_fallback = False                  # si algun evento cayo al mecanismo antiguo
    if reglas and not eventos:
        anotar(f"T08 no declara horas enteras para {planta} (solo minutos sueltos): "
               f"no hay cambio de hora que corregir")
    for fecha, papel, horas in eventos:
        etiqueta = f"{fecha.isoformat()} ({papel})"

        fiable, motivo_dst = es_dst_fiable(fmt, fecha)
        if fiable:
            anotar(f"{etiqueta}: {motivo_dst} -> PVsyst ya corrige este cambio de hora "
                   f"el solo (UsesDST); no se reescribe el CSV para no duplicarlo")
            usesdst_resuelve.append(fecha.isoformat())
            continue

        hubo_fallback = True
        avisar(f"{etiqueta}: UsesDST no es de fiar para resolver este cambio de hora "
               f"({motivo_dst}) -> se cae al mecanismo antiguo de reescribir el CSV a mano. "
               f"Revisa UsesDST/RefTimeS/DateHEte/DateHHiver en la MEF para que PVsyst lo "
               f"resuelva solo y no haga falta este fallback.")
        # TODO(hourshift): aqui es donde engancha el fallback basado en HourShift
        # que queda por definir (pendiente de indicaciones) para las plantas sin
        # UsesDST fiable. De momento el fallback sigue siendo el mecanismo antiguo
        # de reescribir las fechas del CSV (mismo comportamiento que antes de este
        # cambio, ver detectar_cambio_hora / cambio_anterior_al_csv mas abajo).

        cambio = instante_de_cambio(fecha)

        if marcas[0].valor > cambio:
            art = cambio_anterior_al_csv(marcas, fecha, papel, horas)
            anotar(f"{etiqueta}: anterior al CSV -> {art.detalle}")
            anotar(f"            se aplica {art.horas:+d} h a TODO el fichero")
            correcciones.append(art)
            continue

        if marcas[-1].valor < cambio:
            anotar(f"{etiqueta}: posterior al ultimo dato del CSV "
                   f"({marcas[-1].valor:%Y-%m-%d %H:%M}), no le afecta")
            continue

        art, motivo = detectar_cambio_hora(marcas, fecha, fmt["paso"], horas)
        if art is None:
            anotar(f"{etiqueta}: {motivo}")
            avisar(f"{etiqueta}: el cambio de hora cae dentro del CSV pero no se ve en los "
                   f"datos, asi que NO se corrige la hora ({motivo})")
            continue

        if art.tipo != TIPO_ESPERADO[papel]:
            avisar(f"{etiqueta}: T08 lo declara como '{papel}' del tramo, asi que se "
                   f"esperaba {TIPO_ESPERADO[papel]}, pero el CSV trae {art.tipo}. "
                   f"Manda el CSV; revisa si las fechas de T08 estan al reves.")
        anotar(f"{etiqueta}: {art.tipo} -> {art.detalle}")
        anotar(f"            se aplica {art.horas:+d} h desde la linea {art.linea + 1}")
        correcciones.append(art)

    if hubo_fallback and not correcciones:
        avisar("T08 declara horas de desfase, UsesDST no es de fiar para alguna fecha, y el "
               "mecanismo antiguo tampoco ha encontrado un cambio de hora aplicable: NO se "
               "corrige ninguna hora.")

    # --- CSV corregido (copia; el original no se toca) ---
    csv_conv = csv
    hechas: list[Artefacto] = []
    if correcciones:
        destino = Path(tempfile.gettempdir()) / CARPETA_TSFIX / csv.name
        try:
            cambiadas, registro = corregir_csv(csv, lineas, fmt, correcciones, destino)
        except Exception as e:
            return fallar(f"no se pudo escribir el CSV corregido: {e}")
        if cambiadas:
            csv_conv = destino
            hechas = correcciones
            anotar(f"CSV corregido: {cambiadas:,} filas desplazadas")
            for r in registro:
                anotar(f"               {r}")
            anotar(f"               {destino}")
        else:
            # Pasa cuando las correcciones se anulan entre si (p.ej. el tramo se
            # abrio y se cerro antes de que empiecen los datos): el fichero
            # quedaria identico, asi que se convierte el original.
            anotar("las correcciones se anulan entre si: el CSV no cambia")
            try:
                destino.unlink()
            except OSError:
                pass

    # --- minutos sueltos -> -imt ---
    resto, avisos_imt = _resto_imt(reglas, marcas[-1].valor.date())
    for a in avisos_imt:
        avisar(a)
    anotar(f"-imt: {resto:+d} min" + ("" if resto else "  (no se pasa -imt)"))

    # --- salida para el .bat ---
    contenido = f"imt={resto}\ncsv={csv_conv}\n"
    if args.out:
        try:
            Path(args.out).write_text(contenido, encoding="ascii")
        except OSError as e:
            return fallar(f"no se pudo escribir {args.out}: {e}")

    # --- la linea que se ve en el log ---
    trozos = [_resumen_correccion(c) for c in hechas]
    trozos.append(f"-imt:{resto}" if resto else "sin -imt")
    if not hechas:
        if usesdst_resuelve:
            trozos.insert(0, f"UsesDST resuelve la hora ({', '.join(usesdst_resuelve)})")
        else:
            trozos.insert(0, "sin correccion de hora")
    if not args.verbose:
        print(f"{PREF}[TIME-SHIFT] {planta}: {', '.join(trozos)}")
    for a in avisos:
        print(f"{PREF}  [AVISO] {a}")
    if not args.out:
        print(contenido, end="")
    return 0


def cmd_mostrar(args) -> int:
    cfg = cargar_config()
    try:
        xlsx, procedencia = localizar_excel(cfg)
        reglas_todas = leer_reglas(xlsx)
    except Exception as e:
        print(f"[ERROR] {e}")
        return 1

    ws = Path((cfg.get("pvsyst") or {}).get("workspace") or WS_DEFAULT)

    print(f"{NOMBRE_TABLA}: {procedencia}")
    plantas = [p.upper() for p in args.plantas] or sorted(reglas_todas)
    if not plantas:
        print("  (la tabla no tiene ninguna regla)")
        return 0

    for planta in plantas:
        reglas = reglas_todas.get(planta, [])
        print(f"\n{planta}")
        if not reglas:
            print("  (sin reglas)")
            continue

        # Con la MEF a mano se puede decir si UsesDST resuelve cada fecha solo
        # o si haria falta el fallback de reescribir el CSV. Sin ella (MEF no
        # encontrada) se avisa y no se afirma nada sobre el mecanismo.
        mef_path = ws / "Meteo" / f"{planta}.MEF"
        fmt = None
        if mef_path.is_file():
            try:
                fmt = leer_formato_mef(mef_path)
            except Exception:
                fmt = None

        for r in reglas:
            horas, resto = descomponer(r.minutos)
            print(f"  fila {r.fila}: {r}")
            if horas and r.fechas():
                for f in r.fechas():
                    if fmt is None:
                        print(f"     {horas:+d} h en {f.isoformat()} -> no se pudo leer "
                              f"{mef_path} para saber si UsesDST lo resuelve solo")
                        continue
                    fiable, motivo_dst = es_dst_fiable(fmt, f)
                    if fiable:
                        print(f"     {horas:+d} h en {f.isoformat()} -> UsesDST lo resuelve "
                              f"solo ({motivo_dst}); no se toca el CSV")
                    else:
                        print(f"     {horas:+d} h en {f.isoformat()} -> se reescriben las "
                              f"fechas del CSV (UsesDST no es de fiar: {motivo_dst})")
            elif horas:
                print(f"     {horas:+d} h -> (sin fechas: no se corrige)")
            print(f"     {resto:+d} min -> {'convert-meteo -imt:%d' % resto if resto else 'sin -imt'}")
    return 0


def cmd_revisar(args) -> int:
    """Igual que `aplicar` pero sin escribir nada: solo dice que hay en las fechas."""
    args.out = None
    args.dry_run = True
    cfg = cargar_config()
    planta = args.planta.upper()
    ws = Path(args.ws or (cfg.get("pvsyst") or {}).get("workspace") or WS_DEFAULT)
    mef = Path(args.mef) if args.mef else ws / "Meteo" / f"{planta}.MEF"
    csv = Path(args.csv) if args.csv else None

    try:
        xlsx, procedencia = localizar_excel(cfg)
        reglas = leer_reglas(xlsx).get(planta, [])
    except Exception as e:
        print(f"[ERROR] {e}")
        return 1
    if csv is None or not csv.is_file():
        print(f"[ERROR] hace falta --csv y no existe: {csv}")
        return 1
    if not mef.is_file():
        print(f"[ERROR] no existe el MEF {mef}")
        return 1

    fmt = leer_formato_mef(mef)
    lineas = _leer_lineas(csv)
    marcas, ilegibles = leer_marcas(lineas, fmt)
    print(f"{planta}  ({procedencia})")
    print(f"  CSV {csv.name}: {len(marcas):,} marcas, paso declarado {fmt['paso']} min, "
          f"{marcas[0].valor} .. {marcas[-1].valor}")

    pasos = Counter(int((b.valor - a.valor).total_seconds() // 60)
                    for a, b in zip(marcas, marcas[1:]))
    print(f"  pasos observados: {dict(sorted(pasos.items()))}")

    for fecha, papel, horas in eventos_de_cambio(reglas):
        etiqueta = f"{fecha.isoformat()} ({papel})"

        fiable, motivo_dst = es_dst_fiable(fmt, fecha)
        if fiable:
            print(f"  {etiqueta}: UsesDST DE FIAR -> {motivo_dst}")
            print(f"              PVsyst lo resuelve solo; NO se reescribe el CSV")
            continue
        print(f"  {etiqueta}: UsesDST NO es de fiar -> {motivo_dst}")
        print(f"              cae al mecanismo antiguo (reescribir el CSV a mano):")

        cambio = instante_de_cambio(fecha)
        if marcas[0].valor > cambio:
            art = cambio_anterior_al_csv(marcas, fecha, papel, horas)
            print(f"    ANTERIOR al CSV ({art.horas:+d} h a todo el fichero)")
        elif marcas[-1].valor < cambio:
            print(f"    POSTERIOR al CSV, no le afecta")
        else:
            art, motivo = detectar_cambio_hora(marcas, fecha, fmt["paso"], horas)
            if art is None:
                print(f"    SIN cambio de hora -> {motivo}")
            else:
                aviso = "" if art.tipo == TIPO_ESPERADO[papel] else \
                    f"  [AVISO: se esperaba {TIPO_ESPERADO[papel]}]"
                print(f"    {art.tipo} ({art.horas:+d} h) -> {art.detalle}{aviso}")
    if not reglas:
        print("  (la planta no tiene filas en T08)")
    return 0


def _fecha_arg(texto: str) -> dt.date:
    for fmt in ("%Y-%m-%d", "%Y%m%d", "%d/%m/%Y"):
        try:
            return dt.datetime.strptime(texto, fmt).date()
        except ValueError:
            continue
    raise argparse.ArgumentTypeError(f"fecha no valida: {texto} (usa YYYY-MM-DD)")


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(
        description="Corrige los cambios de hora en el CSV raw segun T08_time_shifts "
                    "y devuelve los minutos sueltos para el -imt del convert-meteo.")
    sub = ap.add_subparsers(dest="cmd", required=True)

    a = sub.add_parser("aplicar", help="corrige el CSV y devuelve el -imt")
    a.add_argument("planta")
    a.add_argument("--csv", help="CSV raw que se va a convertir (obligatorio)")
    a.add_argument("--ws", help=f"workspace PVsyst (def. {WS_DEFAULT})")
    a.add_argument("--mef", help="ruta del MEF (def. <ws>\\Meteo\\<PLANTA>.MEF)")
    a.add_argument("--out", help="fichero clave=valor para el .bat (imt=, csv=)")
    a.add_argument("-v", "--verbose", action="store_true",
                   help="escribir todo el detalle (por defecto solo una linea; el "
                        "detalle sale igual si algo falla)")
    a.set_defaults(func=cmd_aplicar)

    m = sub.add_parser("mostrar", help="consulta T08 sin tocar nada")
    m.add_argument("plantas", nargs="*")
    m.set_defaults(func=cmd_mostrar)

    r = sub.add_parser("revisar", help="que hay en las fechas de T08, sin escribir nada")
    r.add_argument("planta")
    r.add_argument("--csv", help="CSV raw a revisar")
    r.add_argument("--ws", help=f"workspace PVsyst (def. {WS_DEFAULT})")
    r.add_argument("--mef", help="ruta del MEF")
    r.set_defaults(func=cmd_revisar)

    args = ap.parse_args(argv)
    return args.func(args)


if __name__ == "__main__":
    sys.exit(main())
